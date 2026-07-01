"""
Prepare per-item research inputs for the enrichment workflow.

Outputs:
  enrichment/inputs/<id>.json   - per-item context for each research agent
  enrichment/notes/             - (empty) where agents write research notes
  links.json                    - {id: {github, live}} confident link fixes for the build step
  prints the JSON array of ids  - paste into the Workflow `args`
"""
import openpyxl, json, re, os, difflib

ROOT = os.path.abspath(".")
SRC = "project.xlsx"

def clean(v):
    if v in (None, ""): return None
    s = str(v).strip(); return s or None
def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", str(s).strip().lower()).strip("-") or "untitled"
def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())

# ---- load repos & build lookup
repos = json.load(open("github_repos.json", encoding="utf-8"))
by_name = {norm(r["name"]): r for r in repos}
by_full = {r["full_name"].lower(): r for r in repos}
by_home = {(r.get("homepage") or "").rstrip("/").lower(): r for r in repos if r.get("homepage")}

def match_repo(name, sheet_gh, sheet_lv):
    # 1) explicit github.com/owner/repo in sheet
    for v in (sheet_gh, sheet_lv):
        if not v: continue
        m = re.search(r"github\.com/([\w.-]+/[\w.-]+)", str(v))
        if m and m.group(1).lower() in by_full:
            return by_full[m.group(1).lower()]
    # 2) github.io/<name> hint
    for v in (sheet_gh, sheet_lv):
        if not v: continue
        m = re.search(r"github\.io/([\w.-]+)", str(v))
        if m and norm(m.group(1)) in by_name:
            return by_name[norm(m.group(1))]
    # 3) homepage match (vercel/netlify)
    for v in (sheet_lv, sheet_gh):
        if not v: continue
        h = str(v).rstrip("/").lower()
        if h in by_home: return by_home[h]
    # 4) exact normalized name
    if norm(name) in by_name: return by_name[norm(name)]
    # 5) very-high fuzzy
    cand = difflib.get_close_matches(norm(name), list(by_name), n=1, cutoff=0.9)
    if cand: return by_name[cand[0]]
    return None

# ---- read sheet projects
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
DONE = {"done", "done - deployed"}
ALIASES = {"name":["project name"],"description":["description"],"status":["status"],
           "skills":["skills used","required skills"],"score":["score","project score"],
           "github":["git hub link","github link"],"live":["live link"],"tags":["tags"],
           "approach":["approach"],"methodology":["methodology"],"challenges":["challenges"],
           "outcomes":["outcomes"],"desired":["desired"]}
def rows_of(sheet):
    rows = list(wb[sheet].iter_rows(values_only=True))
    idx = {str(h).strip().lower(): i for i, h in enumerate(rows[0]) if h not in (None,"")}
    def pick(row, names):
        for n in names:
            if n in idx and idx[n] < len(row): return row[idx[n]]
        return None
    out = []
    for row in rows[1:]:
        nm = clean(pick(row, ALIASES["name"]))
        if not nm: continue
        rec = {k: clean(pick(row, v)) for k, v in ALIASES.items()}
        out.append(rec)
    return out

records = rows_of("Working Projects") + rows_of("Pending Projects")

items, links = [], {}
used = set()
for r in records:
    folder = "done" if (r["status"] or "").strip().lower() in DONE else "pending"
    base = slug(r["name"]); fn = base; n = 2
    while (folder, fn) in used: fn = f"{base}-{n}"; n += 1
    used.add((folder, fn))
    iid = f"{folder}__{fn}"
    repo = match_repo(r["name"], r["github"], r["live"])
    repo_ctx = None
    if repo:
        repo_ctx = {"full_name": repo["full_name"], "html_url": repo["html_url"],
                    "homepage": repo.get("homepage"), "description": repo.get("description"),
                    "language": repo.get("language"), "topics": repo.get("topics") or []}
        gh = repo["html_url"]
        lv = repo.get("homepage") or (r["live"] if str(r["live"] or "").startswith("http") else None)
        links[iid] = {"github": gh, "live": lv}
    items.append({
        "id": iid, "kind": folder, "name": r["name"], "status": r["status"],
        "score": r["score"], "description": r["description"], "skills": r["skills"],
        "tags": r["tags"], "sheet_github": r["github"], "sheet_live": r["live"],
        "repo": repo_ctx,
        "file_path": os.path.join(ROOT, "projects", folder, fn + ".md"),
        "note_path": os.path.join(ROOT, "enrichment", "notes", iid + ".md"),
    })

# ---- ideas (3 rows each: label / title / description)
vals = []
for row in list(wb["Project Ideas"].iter_rows(values_only=True))[1:]:
    for c in row:
        if c not in (None,"") and str(c).strip(): vals.append(str(c).strip()); break
groups = [(vals[i], vals[i+1], vals[i+2]) for i in range(0, len(vals)-2, 3)]
for k, (label, title, desc) in enumerate(groups, 1):
    iid = f"idea__{k:02d}"
    fn = f"{k:02d}-{slug(title)}.md"
    items.append({
        "id": iid, "kind": "idea", "name": title, "status": "Idea / not started",
        "score": None, "description": desc, "skills": None, "tags": None,
        "sheet_github": None, "sheet_live": None, "repo": None,
        "file_path": os.path.join(ROOT, "projects", "ideas", fn),
        "note_path": os.path.join(ROOT, "enrichment", "notes", iid + ".md"),
    })

# ---- write outputs
os.makedirs(os.path.join("enrichment", "inputs"), exist_ok=True)
os.makedirs(os.path.join("enrichment", "notes"), exist_ok=True)
for it in items:
    json.dump(it, open(os.path.join("enrichment", "inputs", it["id"] + ".json"), "w", encoding="utf-8"),
              indent=2, ensure_ascii=False)
json.dump(links, open("links.json", "w", encoding="utf-8"), indent=2, ensure_ascii=False)
json.dump(items, open("items.json", "w", encoding="utf-8"), indent=2, ensure_ascii=False)

matched = sum(1 for it in items if it.get("repo"))
print(f"items={len(items)} (done/pending matched to repo: {matched})  ideas={len(groups)}")
print("IDS=" + json.dumps([it["id"] for it in items]))
